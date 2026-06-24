import io
from datetime import datetime
from flask import Blueprint, render_template, send_file
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from db import get_db_connection

export_bp = Blueprint('export', __name__)


def _build_workbook(sheets: dict) -> Workbook:
    wb = Workbook()
    wb.remove(wb.active)

    header_font = Font(bold=True, color='FFFFFF')
    header_fill = PatternFill(fill_type='solid', fgColor='A99A6F')
    header_align = Alignment(horizontal='center', vertical='center')

    for sheet_name, rows in sheets.items():
        ws = wb.create_sheet(title=sheet_name)
        if not rows:
            ws.append(['No data'])
            continue
        headers = list(rows[0].keys())
        ws.append(headers)
        for col_idx, _ in enumerate(headers, start=1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align
        for row in rows:
            ws.append(list(row))
        for col_idx in range(1, len(headers) + 1):
            col_letter = get_column_letter(col_idx)
            max_len = max(
                len(str(ws.cell(row=r, column=col_idx).value or ''))
                for r in range(1, ws.max_row + 1)
            )
            ws.column_dimensions[col_letter].width = min(max_len + 4, 50)
    return wb


def _send_workbook(wb: Workbook, filename: str):
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf, as_attachment=True, download_name=filename,
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@export_bp.route('/Export')
def export():
    return render_template('Export.html')


_SERVICE_COLS = """
            SELECT mfr.mfr_name AS manufacturer, m.md_name AS model,
                   l.loc_name AS location, s.srv_asset_tag AS gauge_serial_number,
                   us.user_name AS calibrated_by, s.srv_owner AS owner,
                   s.srv_calibration_date, s.srv_cert,
                   s.srv_last_date, s.srv_next_date, s.srv_post_condition
            FROM service s
            LEFT JOIN models m ON s.srv_md_id = m.md_id
            LEFT JOIN Manufacturer mfr ON m.md_mfr_id = mfr.mfr_id
            LEFT JOIN location l ON s.srv_loc_id = l.loc_id
            LEFT JOIN users us ON s.srv_user_id = us.user_id
"""


@export_bp.route('/Export/downloadOld')
def export_downloadOld():
    conn = get_db_connection()
    sheets = {
        'Service': conn.execute(
            _SERVICE_COLS + "WHERE s.srv_post_condition = 'Bad' ORDER BY s.srv_next_date ASC"
        ).fetchall(),
    }
    conn.close()
    filename = f"CRIMan_ExportDamagedUnits_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _send_workbook(_build_workbook(sheets), filename)


@export_bp.route('/Export/downloadService')
def export_downloadService():
    conn = get_db_connection()
    sheets = {
        'Service': conn.execute(
            _SERVICE_COLS + "WHERE s.srv_post_condition != 'Bad' ORDER BY s.srv_next_date ASC"
        ).fetchall(),
    }
    conn.close()
    filename = f"CRIMan_ExportService_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _send_workbook(_build_workbook(sheets), filename)


@export_bp.route('/Export/download')
def export_download():
    conn = get_db_connection()
    sheets = {
        'Models': conn.execute("""
            SELECT m.md_name, m.md_size, m.md_endsize, mfr.mfr_name AS manufacturer
            FROM models m
            LEFT JOIN Manufacturer mfr ON m.md_mfr_id = mfr.mfr_id
        """).fetchall(),
        'Service': conn.execute(
            _SERVICE_COLS + "WHERE s.srv_post_condition != 'Bad' ORDER BY s.srv_next_date ASC"
        ).fetchall(),
    }
    conn.close()
    filename = f"CRIMan_Export_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _send_workbook(_build_workbook(sheets), filename)


@export_bp.route('/Export/downloadBackup')
def export_downloadBackup():
    conn = get_db_connection()
    sheets = {
        'Models': conn.execute("SELECT * FROM models").fetchall(),
        'Service': conn.execute("SELECT * FROM service").fetchall(),
        'Manufacturers': conn.execute("SELECT * FROM Manufacturer").fetchall(),
        'Locations': conn.execute("SELECT * FROM location").fetchall(),
        'Users': conn.execute("SELECT * FROM users").fetchall(),
    }
    conn.close()
    filename = f"CRIMan_Export_Backup{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    return _send_workbook(_build_workbook(sheets), filename)
