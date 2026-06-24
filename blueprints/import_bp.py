import os
from flask import Blueprint, render_template, request, send_file
from openpyxl import load_workbook
from db import get_db_connection

import_bp = Blueprint('import_bp', __name__)

IMPORT_TEMPLATE_PATH = os.path.join(os.path.dirname(os.path.abspath(__file__)), '..', 'CRIMan_Import.xlsx')

_IMPORT_SHEETS = {
    'Manufacturers': {
        'table': 'Manufacturer',
        'columns': ['mfr_id', 'mfr_name'],
    },
    'Models': {
        'table': 'models',
        'columns': ['md_id', 'md_name', 'md_size', 'md_endsize', 'md_mfr_id'],
    },
    'Locations': {
        'table': 'location',
        'columns': ['loc_id', 'loc_name', 'loc_description'],
    },
    'Users': {
        'table': 'users',
        'columns': ['user_id', 'user_name', 'user_loc_id'],
    },
    'Service': {
        'table': 'service',
        'columns': ['srv_id', 'srv_description', 'srv_md_id', 'srv_loc_id',
                    'srv_asset_tag', 'srv_user_id', 'srv_owner', 'srv_calibration_date',
                    'srv_cert', 'srv_last_date', 'srv_next_date', 'srv_post_condition'],
    },
}


@import_bp.route('/Import')
def import_page():
    return render_template('Import.html')


@import_bp.route('/Import/template')
def import_template():
    return send_file(IMPORT_TEMPLATE_PATH, as_attachment=True,
                     download_name='CRIMan_Import.xlsx',
                     mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')


@import_bp.route('/Import/upload', methods=['POST'])
def import_upload():
    if 'file' not in request.files or not request.files['file'].filename:
        return render_template('Import.html', error='No file selected.')

    file = request.files['file']
    if not file.filename.lower().endswith('.xlsx'):
        return render_template('Import.html', error='Please upload an .xlsx file.')

    try:
        wb = load_workbook(file, data_only=True)
    except Exception as e:
        return render_template('Import.html', error=f'Could not read file: {e}')

    conn = get_db_connection()
    counts = {}

    for sheet_name, cfg in _IMPORT_SHEETS.items():
        if sheet_name not in wb.sheetnames:
            counts[sheet_name] = 0
            continue
        ws = wb[sheet_name]
        columns = cfg['columns']
        table = cfg['table']

        headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
        col_indices = {col: headers.index(col) for col in columns if col in headers}

        n = 0
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not any(v is not None for v in row):
                continue
            values = [row[col_indices[col]] if col in col_indices else None for col in columns]
            if values[0] is None:
                continue
            placeholders = ', '.join('?' for _ in columns)
            col_list = ', '.join(columns)
            conn.execute(
                f'INSERT OR REPLACE INTO {table} ({col_list}) VALUES ({placeholders})',
                values
            )
            n += 1
        counts[sheet_name] = n

    conn.commit()
    conn.close()
    return render_template('Import.html', counts=counts)
